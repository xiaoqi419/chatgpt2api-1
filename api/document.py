from __future__ import annotations

import io
import re
from typing import Any

from fastapi import APIRouter, Header, HTTPException, UploadFile
from fastapi.responses import PlainTextResponse

from api.support import require_identity

MAX_FILE_SIZE = 50 * 1024 * 1024  # 50MB


def _try_import(module: str, package: str | None = None) -> Any:
    try:
        return __import__(module)
    except ImportError:
        return None


def _extract_docx(data: bytes) -> str:
    docx = _try_import("docx")
    if docx is None:
        raise HTTPException(status_code=501, detail="DOCX parsing library not installed. Run: uv add python-docx")
    doc = docx.Document(io.BytesIO(data))
    parts: list[str] = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        rows: list[str] = []
        for row in table.rows:
            rows.append(" | ".join(cell.text.strip() for cell in row.cells))
        if rows:
            parts.append("--- 表格 ---\n" + "\n".join(rows))
    return "\n".join(parts)


def _extract_pdf(data: bytes) -> str:
    fitz = _try_import("fitz")
    if fitz is None:
        raise HTTPException(status_code=501, detail="PDF parsing library not installed. Run: uv add PyMuPDF")
    doc = fitz.open(stream=data, filetype="pdf")
    pages: list[str] = []
    for num, page in enumerate(doc, 1):
        text = page.get_text().strip()
        text = re.sub(r"\n{3,}", "\n\n", text)
        if text:
            pages.append(f"--- 第 {num} 页 ---\n{text}")
    doc.close()
    return "\n\n".join(pages)


def _extract_pptx(data: bytes) -> str:
    try:
        from pptx import Presentation
    except ImportError:
        raise HTTPException(status_code=501, detail="PPTX parsing library not installed. Run: uv add python-pptx")
    prs = Presentation(io.BytesIO(data))
    slides: list[str] = []
    for num, slide in enumerate(prs.slides, 1):
        texts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        texts.append(t)
            if shape.has_table:
                for row in shape.table.rows:
                    texts.append(" | ".join(cell.text.strip() for cell in row.cells))
        if texts:
            slides.append(f"--- 第 {num} 页 ---\n" + "\n".join(texts))
    return "\n\n".join(slides)


def _read_text(data: bytes) -> str:
    for enc in ("utf-8", "gbk", "gb2312", "gb18030", "latin-1"):
        try:
            return data.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("utf-8", errors="replace")


SUPPORTED: dict[str, tuple[str, str, Any]] = {
    ".docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "Word 文档", _extract_docx),
    ".pdf": ("application/pdf", "PDF 文档", _extract_pdf),
    ".pptx": ("application/vnd.openxmlformats-officedocument.presentationml.presentation", "PPT 文档", _extract_pptx),
    ".txt": ("text/plain", "纯文本", lambda d: _read_text(d)),
    ".md": ("text/markdown", "Markdown", lambda d: _read_text(d)),
    ".json": ("application/json", "JSON", lambda d: _read_text(d)),
    ".csv": ("text/csv", "CSV", lambda d: _read_text(d)),
    ".yaml": ("application/x-yaml", "YAML", lambda d: _read_text(d)),
    ".yml": ("application/x-yaml", "YAML", lambda d: _read_text(d)),
    ".xml": ("application/xml", "XML", lambda d: _read_text(d)),
    ".log": ("text/plain", "日志", lambda d: _read_text(d)),
    ".toml": ("text/plain", "TOML", lambda d: _read_text(d)),
    ".ini": ("text/plain", "INI 配置", lambda d: _read_text(d)),
    ".cfg": ("text/plain", "配置文件", lambda d: _read_text(d)),
    ".conf": ("text/plain", "配置文件", lambda d: _read_text(d)),
    ".sh": ("text/x-shellscript", "Shell 脚本", lambda d: _read_text(d)),
    ".bat": ("text/x-bat", "批处理", lambda d: _read_text(d)),
    ".py": ("text/x-python", "Python", lambda d: _read_text(d)),
    ".js": ("application/javascript", "JavaScript", lambda d: _read_text(d)),
    ".ts": ("application/typescript", "TypeScript", lambda d: _read_text(d)),
    ".html": ("text/html", "HTML", lambda d: _read_text(d)),
    ".css": ("text/css", "CSS", lambda d: _read_text(d)),
    ".sql": ("text/x-sql", "SQL", lambda d: _read_text(d)),
    ".env": ("text/plain", "环境变量", lambda d: _read_text(d)),
    ".gitignore": ("text/plain", "Git 忽略规则", lambda d: _read_text(d)),
}


def create_router() -> APIRouter:
    router = APIRouter(prefix="/api")

    @router.post("/extract-text")
    async def extract_text(
        file: UploadFile,
        authorization: str | None = Header(default=None),
    ):
        """上传文档并提取文本内容。支持 DOCX、PDF、PPTX、TXT、代码文件等。"""
        require_identity(authorization)

        filename = file.filename or "untitled"
        ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

        if ext not in SUPPORTED:
            supported_list = ", ".join(sorted(SUPPORTED.keys()))
            raise HTTPException(
                status_code=400,
                detail=f"不支持的文件格式: {ext}。支持的格式: {supported_list}",
            )

        try:
            data = await file.read()
        except Exception as exc:
            raise HTTPException(status_code=400, detail=f"读取文件失败: {exc}") from exc

        if not data:
            raise HTTPException(status_code=400, detail="文件内容为空")
        if len(data) > MAX_FILE_SIZE:
            raise HTTPException(status_code=400, detail="文件超过 50MB 限制")

        mime, label, extractor = SUPPORTED[ext]
        try:
            text = extractor(data)
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(status_code=500, detail=f"提取文本失败: {exc}") from exc

        if not text.strip():
            text = "(空内容)"

        return {"filename": filename, "type": label, "size": len(data), "text": text}

    return router
